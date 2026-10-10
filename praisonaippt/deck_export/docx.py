"""Export deck dicts to Word (.docx) with underline + highlight on YAML phrases."""

from __future__ import annotations

import re
from pathlib import Path
from typing import Any, List, Optional, Tuple

from docx import Document
from docx.enum.text import WD_COLOR_INDEX, WD_UNDERLINE, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt

DEFAULT_ENGLISH_FONT = "Palatino Linotype"
DEFAULT_BODY_SIZE_PT = 12

from ..slide_renderers import resolve_renderer
from .highlights import deck_default_highlight_hex, normalise_phrases, resolve_highlight_hex
from .markdown import _render_verse_entry, _section_heading, _strip


def _colour_index(colour_hint: object | None, default_hex: str) -> WD_COLOR_INDEX:
    """Map YAML highlight colour to Word's limited highlight palette."""
    if colour_hint is None or colour_hint == "":
        hex_color = default_hex.upper()
    else:
        hex_color = resolve_highlight_hex(colour_hint, default_hex).upper()
    named = str(colour_hint or "").strip().lower()
    if named == "green" or hex_color in ("#32B432", "#00FF00", "#008000"):
        return WD_COLOR_INDEX.BRIGHT_GREEN
    if named in ("red",) or hex_color.startswith("#DC") or hex_color.startswith("#FF0000"):
        return WD_COLOR_INDEX.RED
    if named in ("blue", "cyan") or hex_color.startswith("#1E"):
        return WD_COLOR_INDEX.TURQUOISE
    if named == "purple":
        return WD_COLOR_INDEX.VIOLET
    return WD_COLOR_INDEX.YELLOW


def _highlight_spans(
    text: str,
    highlights: list | None,
    *,
    default_hex: str,
) -> List[Tuple[str, Optional[WD_COLOR_INDEX]]]:
    """Split *text* into (segment, highlight_index) tuples; ``None`` = plain."""
    if not text:
        return []
    items = normalise_phrases(highlights)
    if not items:
        return [(text, None)]

    matches: List[Tuple[int, int, WD_COLOR_INDEX]] = []
    for phrase, colour_hint in items:
        if not phrase:
            continue
        idx = _colour_index(colour_hint, default_hex)
        for match in re.finditer(re.escape(phrase), text, re.IGNORECASE):
            matches.append((match.start(), match.end(), idx))

    if not matches:
        return [(text, None)]

    matches.sort(key=lambda m: (m[0], -(m[1] - m[0])))
    occupied: List[Tuple[int, int]] = []
    chosen: List[Tuple[int, int, WD_COLOR_INDEX]] = []
    for start, end, idx in matches:
        if any(not (end <= a or start >= b) for a, b in occupied):
            continue
        occupied.append((start, end))
        chosen.append((start, end, idx))
    chosen.sort(key=lambda m: m[0])

    out: List[Tuple[str, Optional[WD_COLOR_INDEX]]] = []
    pos = 0
    for start, end, idx in chosen:
        if start > pos:
            out.append((text[pos:start], None))
        out.append((text[start:end], idx))
        pos = end
    if pos < len(text):
        out.append((text[pos:], None))
    return out


def _set_run_font(run, face: str, *, complex_script: bool = False) -> None:
    """Apply *face* to Latin and complex-script slots (Tamil needs ``w:cs``)."""
    run.font.name = face
    r_pr = run._element.get_or_add_rPr()
    r_fonts = r_pr.find(qn("w:rFonts"))
    if r_fonts is None:
        r_fonts = OxmlElement("w:rFonts")
        r_pr.insert(0, r_fonts)
    r_fonts.set(qn("w:ascii"), face)
    r_fonts.set(qn("w:hAnsi"), face)
    if complex_script:
        r_fonts.set(qn("w:cs"), face)
        r_fonts.set(qn("w:eastAsia"), face)


def _apply_document_layout(doc: Document) -> None:
    """Print-friendly margins and body typography."""
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    normal = doc.styles["Normal"]
    normal.font.name = DEFAULT_ENGLISH_FONT
    normal.font.size = Pt(DEFAULT_BODY_SIZE_PT)
    pf = normal.paragraph_format
    pf.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    pf.line_spacing = 1.15
    pf.space_after = Pt(6)

    quote = doc.styles["Intense Quote"]
    quote.font.name = DEFAULT_ENGLISH_FONT
    quote.font.size = Pt(DEFAULT_BODY_SIZE_PT)
    quote.paragraph_format.line_spacing = 1.15
    quote.paragraph_format.space_before = Pt(6)
    quote.paragraph_format.space_after = Pt(6)
    quote.paragraph_format.left_indent = Inches(0.35)
    quote.paragraph_format.right_indent = Inches(0.15)


def _add_section_rule(doc: Document) -> None:
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    p.paragraph_format.space_after = Pt(14)
    p_pr = p._element.get_or_add_pPr()
    p_bdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "A0A0A0")
    p_bdr.append(bottom)
    p_pr.append(p_bdr)


def _add_runs(
    paragraph,
    text: str,
    highlights: list | None,
    *,
    default_hex: str,
    font_size: int = DEFAULT_BODY_SIZE_PT,
    font_name: str | None = DEFAULT_ENGLISH_FONT,
    complex_script: bool = False,
) -> None:
    for segment, hl_idx in _highlight_spans(text, highlights, default_hex=default_hex):
        run = paragraph.add_run(segment)
        run.font.size = Pt(font_size)
        if font_name:
            _set_run_font(run, font_name, complex_script=complex_script)
        if hl_idx is not None:
            run.font.bold = True
            run.font.underline = WD_UNDERLINE.SINGLE
            run.font.highlight_color = hl_idx


def _add_quote_block(
    doc: Document,
    body: str,
    reference: str,
    highlights: list | None,
    *,
    default_hex: str,
    font_name: str | None = None,
    label: str | None = None,
) -> None:
    if label:
        lp = doc.add_paragraph()
        lr = lp.add_run(label)
        lr.font.size = Pt(10)
        lr.font.bold = True
        lr.font.italic = True
    if body:
        p = doc.add_paragraph(style="Intense Quote")
        _add_runs(
            p,
            f"\u201c{body}\u201d",
            highlights,
            default_hex=default_hex,
            font_name=font_name or DEFAULT_ENGLISH_FONT,
            complex_script=bool(font_name),
        )
    ref = _strip(reference)
    if ref:
        rp = doc.add_paragraph()
        rp.paragraph_format.left_indent = Inches(0.35)
        rr = rp.add_run(f"\u2014 {ref}")
        rr.font.size = Pt(11)
        rr.font.italic = True
        if font_name:
            _set_run_font(rr, font_name, complex_script=True)
        else:
            _set_run_font(rr, DEFAULT_ENGLISH_FONT)


def _append_verse_docx(
    doc: Document,
    verse: dict,
    *,
    highlights: bool,
    default_hex: str,
    tamil_font: str | None = None,
    bilingual: bool = False,
    tamil_first: bool = False,
) -> None:
    if not isinstance(verse, dict):
        return
    try:
        kind = resolve_renderer(verse).kind
    except Exception:
        kind = _strip(verse.get("slide_type")) or "verse"

    if kind in ("verse", "hebrew_rename", "quote", "avatar_quote") or (
        _strip(verse.get("text")) or _strip(verse.get("reference"))
    ):
        text = _strip(verse.get("text"))
        ref = _strip(verse.get("reference"))
        below = _strip(verse.get("text_below_reference"))
        leading = _strip(verse.get("leading_title"))
        leading_ta = _strip(verse.get("leading_title_ta"))
        if leading_ta or leading:
            title = " / ".join(x for x in (leading_ta, leading) if x) if bilingual else (leading or leading_ta)
            doc.add_heading(title, level=3)
        hl = verse.get("highlights") if highlights else None
        hl_ta = verse.get("highlights_ta") if highlights else None
        text_ta = _strip(verse.get("text_ta"))
        ref_ta = _strip(verse.get("reference_ta"))

        def _en_block() -> None:
            if text:
                _add_quote_block(doc, text, ref, hl, default_hex=default_hex)
            elif ref:
                doc.add_heading(ref, level=3)

        def _ta_block() -> None:
            if text_ta:
                _add_quote_block(
                    doc,
                    text_ta,
                    ref_ta,
                    hl_ta,
                    default_hex=default_hex,
                    font_name=tamil_font,
                    label="TAOVBSI" if bilingual else None,
                )

        if bilingual and text_ta:
            if tamil_first:
                _ta_block()
                _en_block()
            else:
                _en_block()
                _ta_block()
        else:
            _en_block()
        if below:
            hl2 = verse.get("text_below_reference_highlights") if highlights else None
            _add_quote_block(doc, below, "", hl2, default_hex=default_hex)
        return

    md = _render_verse_entry(verse, highlights=highlights)
    if md:
        for block in md.split("\n\n"):
            block = block.strip()
            if not block:
                continue
            if block.startswith("### "):
                doc.add_heading(block[4:].strip(), level=3)
            elif block.startswith("> "):
                lines = block.splitlines()
                body = lines[0][2:].strip().strip('"')
                ref_line = ""
                if len(lines) > 1 and lines[1].startswith("> — "):
                    ref_line = lines[1][4:].strip()
                plain = re.sub(r"\*\*(.+?)\*\*", r"\1", body)
                _add_quote_block(doc, plain, ref_line, None, default_hex=default_hex)
            else:
                doc.add_paragraph(block)


def deck_wants_bilingual_export(data: dict) -> bool:
    from ..bilingual import BilingualConfig

    if BilingualConfig.from_deck(data).enabled:
        return True
    for section in data.get("sections") or []:
        if not isinstance(section, dict):
            continue
        if _strip(section.get("section_ta")):
            return True
        for verse in section.get("verses") or []:
            if not isinstance(verse, dict):
                continue
            if _strip(verse.get("text_ta")) or _strip(verse.get("reference_ta")):
                return True
            if verse.get("highlights_ta"):
                return True
    return False


def _docx_export_data(data: dict) -> dict:
    """Fill Tamil scripture text when the deck carries bilingual metadata."""
    if not deck_wants_bilingual_export(data):
        return data
    from ..bilingual import BilingualConfig, prepare_bilingual_deck

    work = dict(data)
    if not BilingualConfig.from_deck(work).enabled:
        work["bilingual"] = {"language": "ta", "version": "TAOVBSI", "strict": False}
    filled, _report = prepare_bilingual_deck(work)
    return filled


def deck_to_docx(
    data: dict,
    *,
    highlights: bool = True,
    separators: bool = True,
    include_title: bool = True,
) -> Document:
    """Build a ``python-docx`` document mirroring deck markdown structure."""
    from ..bilingual_slides import bilingual_primary, tamil_font

    data = _docx_export_data(data)
    style = data.get("slide_style") or {}
    default_hex = deck_default_highlight_hex(style)
    ta_face = tamil_font(style)
    bilingual = deck_wants_bilingual_export(data)
    tamil_first = bilingual_primary(style) == "ta"
    doc = Document()
    _apply_document_layout(doc)

    if include_title:
        title = _strip(data.get("presentation_title"))
        if title:
            doc.add_heading(title, level=0)
        subtitle = _strip(data.get("presentation_subtitle"))
        if subtitle:
            doc.add_paragraph(subtitle)

    sections = data.get("sections") or []
    prev_had_verses = False

    for section in sections:
        if not isinstance(section, dict):
            continue
        section_name = _strip(section.get("section"))
        verses = section.get("verses") or []

        if separators and prev_had_verses:
            _add_section_rule(doc)

        section_ta = _strip(section.get("section_ta"))
        section_sub = _strip(section.get("section_subtitle"))
        section_sub_ta = _strip(section.get("section_subtitle_ta"))
        if section_name or section_ta:
            if bilingual and section_ta and section_name:
                doc.add_heading(f"{section_ta} / {section_name}", level=2)
            elif section_ta:
                doc.add_heading(section_ta, level=2)
            else:
                heading_md = _section_heading(section_name, verses)
                if heading_md.startswith("## "):
                    doc.add_heading(section_name, level=2)
                elif heading_md.startswith("# "):
                    doc.add_heading(section_name, level=1)
            if section_sub or section_sub_ta:
                sub = " / ".join(x for x in (section_sub_ta, section_sub) if x)
                sp = doc.add_paragraph(sub)
                sp.runs[0].font.italic = True

        had = False
        for verse in verses:
            before = len(doc.paragraphs)
            _append_verse_docx(
                doc,
                verse,
                highlights=highlights,
                default_hex=default_hex,
                tamil_font=ta_face,
                bilingual=bilingual,
                tamil_first=tamil_first,
            )
            if len(doc.paragraphs) > before:
                had = True
        prev_had_verses = had

    return doc


def write_deck_docx(
    data: dict,
    output_path: str | Path,
    **kwargs: Any,
) -> Path:
    path = Path(output_path)
    doc = deck_to_docx(data, **kwargs)
    doc.save(str(path))
    return path
