"""Render English + Tamil scripture slides.

``slide_style.bilingual_layout``:

* ``cards`` (default) — two rounded translucent cards, one per language, sized
  from measured text. A verse that does not fit at the minimum size is split
  at clause breaks into "(1/2)", "(2/2)" slides, both languages in step.
* ``paired`` — one full slide per language, back to back.
* ``below`` — the original layout (Tamil block under the English); kept for
  existing decks.

``slide_style.bilingual_primary`` (``english`` | ``tamil``) decides which
language comes first. ``bilingual_overflow: paired`` renders a verse that would
need splitting as paired slides instead.
"""

from __future__ import annotations

import copy
import logging
from dataclasses import dataclass
from itertools import zip_longest
from typing import List, Optional, Tuple

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt

from .layout_tokens import layout_in, typography_pt
from .text_measure import TAMIL_RE, TextMeasurer, split_balanced

logger = logging.getLogger(__name__)

LAYOUTS = ("cards", "paired", "below")
OVERFLOW = ("split", "paired")
_PRIMARY = {"english": "en", "en": "en", "tamil": "ta", "ta": "ta"}
DEFAULT_TAMIL_FONT = "Noto Serif Tamil"
_KIND = "bilingual_cards"


def bilingual_layout(style: dict) -> str:
    return str((style or {}).get("bilingual_layout") or "cards").strip().lower()


def bilingual_primary(style: dict) -> str:
    return _PRIMARY.get(str((style or {}).get("bilingual_primary") or "english").strip().lower(), "en")


def tamil_font(style: dict) -> str:
    return (style or {}).get("tamil_font_name") or DEFAULT_TAMIL_FONT


# ------------------------------------------------------------------ Tamil fonts


def set_tamil_run_font(run, face: str, lang: str = "ta-IN") -> None:
    """Point every script slot at the Tamil face.

    python-pptx's ``font.name`` only writes ``<a:latin>``; PowerPoint draws
    Tamil from ``<a:cs>`` (complex script) and otherwise falls back to the
    theme's Latha, so the chosen face never shows. Google Slides and Keynote
    read ``<a:latin>``. Setting all three plus ``lang`` works everywhere.
    """
    run.font.name = face
    rPr = run._r.get_or_add_rPr()
    rPr.set("lang", lang)
    for tag in ("a:ea", "a:cs"):
        old = rPr.find(qn(tag))
        if old is not None:
            rPr.remove(old)
    latin = rPr.find(qn("a:latin"))
    ea = etree.SubElement(rPr, qn("a:ea"))
    ea.set("typeface", face)
    cs = etree.SubElement(rPr, qn("a:cs"))
    cs.set("typeface", face)
    latin.addnext(ea)
    ea.addnext(cs)


def _script_segments(text: str) -> List[Tuple[str, bool]]:
    """Split text into (segment, is_tamil); spaces/punctuation stay with their neighbour."""
    segs: List[List] = []
    for ch in text:
        if TAMIL_RE.match(ch):
            kind = True
        elif ch.isalpha():
            kind = False
        else:
            kind = None
        if segs and (kind is None or segs[-1][1] is None or segs[-1][1] == kind):
            segs[-1][0] += ch
            if segs[-1][1] is None:
                segs[-1][1] = kind
        else:
            segs.append([ch, kind])
    return [(s, bool(k)) for s, k in segs]


def _iter_text_frames(shapes):
    for shape in shapes:
        if shape.shape_type == 6:  # group
            yield from _iter_text_frames(shape.shapes)
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            yield shape.text_frame
        if getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                for cell in row.cells:
                    yield cell.text_frame


def apply_tamil_fonts(prs, face: str) -> int:
    """Give every Tamil run the Tamil face in all script slots; split mixed runs first.

    Returns the number of Tamil runs touched. Covers title, section and legacy
    slides too, e.g. "He will not put to shame / அவன் வெட்கப்படுவதில்லை".
    """
    from pptx.text.text import _Run

    touched = 0
    for slide in prs.slides:
        for tf in _iter_text_frames(slide.shapes):
            for para in tf.paragraphs:
                for run in list(para.runs):
                    text = run.text
                    if not TAMIL_RE.search(text):
                        continue
                    segs = _script_segments(text)
                    if len(segs) == 1:
                        set_tamil_run_font(run, face)
                        touched += 1
                        continue
                    r = run._r
                    r.find(qn("a:t")).text = segs[0][0]
                    prev = r
                    elems = [(r, segs[0][1])]
                    for seg, is_ta in segs[1:]:
                        new = copy.deepcopy(r)
                        new.find(qn("a:t")).text = seg
                        prev.addnext(new)
                        prev = new
                        elems.append((new, is_ta))
                    for el, is_ta in elems:
                        if is_ta:
                            set_tamil_run_font(_Run(el, para), face)
                            touched += 1
    return touched


# ------------------------------------------------------------------ planning


@dataclass
class Block:
    lang: str
    reference: str
    text: str
    highlights: list
    size: int = 0


@dataclass
class SlidePlan:
    blocks: List[Block]
    title: Optional[List[Tuple[str, str]]]   # [(text, lang)] or None
    en_size: int
    ta_size: int


class CardsGeometry:
    def __init__(self, prs, style: dict):
        g = lambda key: float(layout_in(style, _KIND, key))  # noqa: E731
        self.slide_w = prs.slide_width.inches
        self.slide_h = prs.slide_height.inches
        self.card_w = min(self.slide_w - 2 * g("margin_in"), g("max_width_in"))
        self.left = (self.slide_w - self.card_w) / 2.0
        self.top = g("top_in")
        self.bottom = g("bottom_in")
        self.gap = g("card_gap_in")
        self.pad_x = g("pad_x_in")
        self.pad_top = g("pad_top_in")
        self.pad_bottom = g("pad_bottom_in")
        self.ref_gap = g("ref_gap_in")
        self.title_gap = g("title_gap_in")
        self.corner = g("corner_in")
        self.inner_w_pt = (self.card_w - 2 * self.pad_x) * 72.0 - 4.0   # small safety for PowerPoint's own wrap


class CardsPlanner:
    def __init__(self, prs, style: dict):
        self.style = style
        self.geo = CardsGeometry(prs, style)
        self.en = TextMeasurer(style.get("font_name") or "Palatino", "latin")
        self.ta = TextMeasurer(tamil_font(style), "tamil")
        self.en0 = int(typography_pt(style, "bilingual_en_size_pt"))
        self.ta0 = int(typography_pt(style, "bilingual_ta_size_pt"))
        self.min_pt = int(typography_pt(style, "bilingual_min_size_pt"))
        self.ref_pt = int(typography_pt(style, "bilingual_ref_size_pt"))
        self.title_pt = int(typography_pt(style, "bilingual_title_size_pt"))
        self.max_parts = int(layout_in(style, _KIND, "max_parts"))

    def measurer(self, lang: str) -> TextMeasurer:
        return self.ta if lang == "ta" else self.en

    def title_height_in(self, title) -> float:
        if not title:
            return 0.0
        width = self.geo.card_w * 72.0
        lines = 0
        for text, lang in title:
            lines += len(self.measurer(lang).wrap(text, width, self.title_pt))
        line_pt = max(self.measurer(lang).line_height_pt(self.title_pt) for _, lang in title)
        return lines * line_pt / 72.0 + self.geo.title_gap

    def card_height_in(self, block: Block, size: int) -> float:
        m = self.measurer(block.lang)
        ref_h = m.line_height_pt(self.ref_pt) / 72.0
        body_h = m.height_pt(block.text, self.geo.inner_w_pt, size) / 72.0
        return self.geo.pad_top + ref_h + self.geo.ref_gap + body_h + self.geo.pad_bottom

    def fits(self, blocks: List[Block], title, en: int, ta: int) -> bool:
        avail = self.geo.slide_h - self.geo.top - self.geo.bottom - self.title_height_in(title)
        used = sum(self.card_height_in(b, ta if b.lang == "ta" else en) for b in blocks if b.text)
        used += self.geo.gap * (sum(1 for b in blocks if b.text) - 1)
        return used <= avail + 1e-6

    def size_steps(self):
        d = 0
        while True:
            en, ta = max(self.min_pt, self.en0 - d), max(self.min_pt, self.ta0 - d)
            yield en, ta
            if en == self.min_pt and ta == self.min_pt:
                return
            d += 1

    def plan(self, blocks: List[Block], title) -> Tuple[List[SlidePlan], bool]:
        """Slides for one verse; the bool is False when even max_parts could not fit."""
        for n in range(1, self.max_parts + 1):
            if n == 1:
                parts = [[copy.copy(b) for b in blocks]]
            else:
                pieces = []
                for b in blocks:
                    keep = [h if isinstance(h, str) else (h or {}).get("text", "") for h in b.highlights or []]
                    pieces.append(split_balanced(b.text, n, keep))
                parts = []
                for texts in zip_longest(*pieces, fillvalue=""):
                    parts.append([Block(b.lang, b.reference, t, b.highlights) for b, t in zip(blocks, texts)])
            for en, ta in self.size_steps():
                if all(self.fits(p, title if i == 0 else None, en, ta) for i, p in enumerate(parts)):
                    return self._label(parts, title, en, ta), True
        return self._label(parts, title, self.min_pt, self.min_pt), False

    @staticmethod
    def _label(parts, title, en, ta) -> List[SlidePlan]:
        n = len(parts)
        out = []
        for i, blocks in enumerate(parts):
            if n > 1:
                for b in blocks:
                    b.reference = f"{b.reference} ({i + 1}/{n})" if b.reference else f"({i + 1}/{n})"
            out.append(SlidePlan(blocks, title if i == 0 else None, en, ta))
        return out


# ------------------------------------------------------------------ drawing


def _rgb(hex_str: str) -> RGBColor:
    return RGBColor.from_string(str(hex_str).lstrip("#").upper())


def _set_alpha(color_parent, opacity: float) -> None:
    clr = color_parent.find(qn("a:srgbClr"))
    if clr is None:
        return
    for old in clr.findall(qn("a:alpha")):
        clr.remove(old)
    alpha = etree.SubElement(clr, qn("a:alpha"))
    alpha.set("val", str(int(round(max(0.0, min(1.0, opacity)) * 100000))))


def _card(slide, geo: CardsGeometry, top: float, height: float, fill_hex: str, opacity: float,
          outline_opacity: float):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE, Inches(geo.left), Inches(top), Inches(geo.card_w), Inches(height)
    )
    shape.shadow.inherit = False
    shape.adjustments[0] = min(0.5, geo.corner / max(min(geo.card_w, height), 0.01))
    shape.fill.solid()
    shape.fill.fore_color.rgb = _rgb(fill_hex)
    spPr = shape._element.spPr
    _set_alpha(spPr.find(qn("a:solidFill")), opacity)
    shape.line.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    shape.line.width = Pt(0.75)
    ln = spPr.find(qn("a:ln"))
    if ln is not None:
        _set_alpha(ln.find(qn("a:solidFill")), outline_opacity)
    shape.text_frame.text = ""
    return shape


def _textbox(slide, left, top, width, height):
    tb = slide.shapes.add_textbox(Inches(left), Inches(top), Inches(width), Inches(height))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.auto_size = MSO_AUTO_SIZE.NONE
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = Emu(0)
    return tf


def _draw_cards_slide(prs, plan: SlidePlan, style: dict, planner: CardsPlanner):
    from .core import _apply_highlights, _apply_slide_background, _resolve_theme

    theme = _resolve_theme(style)
    geo = planner.geo
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _apply_slide_background(slide, style, prs)
    tok = lambda key: layout_in(style, _KIND, key)  # noqa: E731
    en_font = theme["font_name"]
    ta_face = tamil_font(style)
    ann_pt = int(typography_pt(style, "annotation_size_pt", 46))

    y = geo.top
    if plan.title:
        title_h = planner.title_height_in(plan.title) - geo.title_gap
        tf = _textbox(slide, geo.left, y, geo.card_w, title_h + 0.1)
        for i, (text, lang) in enumerate(plan.title):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            run = p.add_run()
            run.text = text
            run.font.size = Pt(planner.title_pt)
            run.font.bold = True
            run.font.color.rgb = theme["title"]
            if lang == "ta":
                set_tamil_run_font(run, ta_face)
            elif en_font:
                run.font.name = en_font
        y += title_h + geo.title_gap

    blocks = [b for b in plan.blocks if b.text]
    heights = [planner.card_height_in(b, plan.ta_size if b.lang == "ta" else plan.en_size) for b in blocks]
    spare = geo.slide_h - geo.bottom - y - sum(heights) - geo.gap * (len(blocks) - 1)
    extra = max(spare, 0.0) / max(len(blocks), 1)
    opacities = [float(tok("card_opacity_top")), float(tok("card_opacity_bottom"))]
    ref_rgb = _rgb(tok("reference_color"))
    ta_rgb = _rgb(tok("tamil_text_color"))

    for idx, (block, h) in enumerate(zip(blocks, heights)):
        card_h = h + extra
        opacity = opacities[min(idx, 1)] if len(blocks) > 1 else opacities[0]
        _card(slide, geo, y, card_h, tok("card_fill"), opacity, float(tok("card_outline_opacity")))
        tf = _textbox(
            slide, geo.left + geo.pad_x, y + geo.pad_top,
            geo.card_w - 2 * geo.pad_x, card_h - geo.pad_top - geo.pad_bottom,
        )
        is_ta = block.lang == "ta"
        face = ta_face if is_ta else en_font
        size = plan.ta_size if is_ta else plan.en_size
        rp = tf.paragraphs[0]
        rp.alignment = PP_ALIGN.LEFT
        rp.space_after = Pt(geo.ref_gap * 72)
        rr = rp.add_run()
        rr.text = block.reference or ""
        rr.font.size = Pt(planner.ref_pt)
        rr.font.bold = True
        rr.font.color.rgb = ref_rgb
        if is_ta:
            set_tamil_run_font(rr, face)
        elif face:
            rr.font.name = face
        body_rgb = ta_rgb if is_ta else theme["body"]
        for line in [ln for ln in block.text.split("\n") if ln.strip()] or [""]:
            p = tf.add_paragraph()
            p.alignment = PP_ALIGN.LEFT
            _apply_highlights(
                p, line, block.highlights, None,
                body_rgb=body_rgb, highlight_rgb=theme["highlight"], annotation_rgb=theme["annotation"],
                font_name=face, base_font_size=size, annotation_size_pt=ann_pt,
            )
            if is_ta:
                for run in p.runs:
                    set_tamil_run_font(run, face)
        y += card_h + geo.gap
    return slide


# ------------------------------------------------------------------ entry points


def verse_blocks(verse: dict, style: dict) -> List[Block]:
    en = Block("en", str(verse.get("reference") or ""), str(verse.get("text") or "").strip(),
               list(verse.get("highlights") or []))
    ta = Block("ta", str(verse.get("reference_ta") or ""), str(verse.get("text_ta") or "").strip(),
               list(verse.get("highlights_ta") or []))
    return [ta, en] if bilingual_primary(style) == "ta" else [en, ta]


def verse_title(verse: dict, style: dict):
    en = (verse.get("leading_title") or "").strip()
    ta = (verse.get("leading_title_ta") or "").strip()
    items = [(en, "en"), (ta, "ta")]
    if bilingual_primary(style) == "ta":
        items.reverse()
    items = [(t, lang) for t, lang in items if t]
    return items or None


def plan_cards(prs, verse: dict, style: dict) -> Tuple[List[SlidePlan], bool]:
    planner = CardsPlanner(prs, style)
    return planner.plan(verse_blocks(verse, style), verse_title(verse, style))


def _paired_verses(verse: dict, style: dict) -> List[dict]:
    keep = {k: v for k, v in verse.items() if not k.endswith("_ta") and k != "notes"}
    en = dict(keep)
    ta = dict(keep)
    ta.update(
        reference=verse.get("reference_ta") or verse.get("reference"),
        text=verse.get("text_ta"),
        highlights=verse.get("highlights_ta") or [],
        leading_title=verse.get("leading_title_ta") or verse.get("leading_title"),
        large_text=None,
    )
    ta["slide_style"] = {**(verse.get("slide_style") or {}), "font_name": tamil_font(style)}
    pair = [ta, en] if bilingual_primary(style) == "ta" else [en, ta]
    if verse.get("notes"):
        pair[0]["notes"] = verse["notes"]
    return pair


def render_bilingual_verse(prs, verse: dict, style: dict, *, source_file: Optional[str] = None) -> None:
    from .slide_renderers import VerseRenderer, _apply_notes

    layout = bilingual_layout(style)
    overflow = str(style.get("bilingual_overflow") or "split").lower()
    if layout == "cards":
        planner = CardsPlanner(prs, style)
        plans, fitted = planner.plan(verse_blocks(verse, style), verse_title(verse, style))
        if len(plans) == 1 or overflow != "paired":
            if not fitted:
                logger.warning("%s does not fit in %d slides even at %d pt; text may overflow",
                               verse.get("reference"), len(plans), planner.min_pt)
            for i, plan in enumerate(plans):
                slide = _draw_cards_slide(prs, plan, style, planner)
                if i == 0:
                    _apply_notes(slide, verse)
            return
    renderer = VerseRenderer()
    for single in _paired_verses(verse, style):
        renderer.render(prs, single, style, source_file=source_file)


def legacy_part_count(verse: dict, style: dict) -> int:
    """Slides the single-language verse renderer makes for ``verse``."""
    from .layout_tokens import split_max_length_default
    from .text_measure import is_tamil
    from .utils import split_long_text

    text = verse.get("text", "") or ""
    max_len = int(verse.get("split_max_length") or split_max_length_default(style))
    return len(split_long_text(text, max_length=max(max_len, 50), clause_fallback=is_tamil(text)))


def bilingual_slide_count(verse: dict, style: dict, slide_w_in: float, slide_h_in: float) -> int:
    """How many slides ``render_bilingual_verse`` makes, without building them (video plans)."""

    class _Size:
        slide_width = Inches(slide_w_in)
        slide_height = Inches(slide_h_in)

    if bilingual_layout(style) == "cards":
        planner = CardsPlanner(_Size, style)
        plans, _ = planner.plan(verse_blocks(verse, style), verse_title(verse, style))
        if len(plans) == 1 or str(style.get("bilingual_overflow") or "split").lower() != "paired":
            return len(plans)
    return sum(legacy_part_count(v, style) for v in _paired_verses(verse, style))
