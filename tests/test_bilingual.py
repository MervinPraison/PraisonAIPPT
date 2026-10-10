"""English + Tamil scripture slides: Tamil lookup, fill precedence, layouts, fonts."""

from __future__ import annotations

import copy
import json
from argparse import Namespace

import pytest
from pptx import Presentation
from pptx.oxml.ns import qn

from praisonaippt.bible_books import BookNotFound, lookup_book
from praisonaippt.bilingual import (
    BibleLookupError,
    BilingualConfig,
    fill_tamil,
    freeze_bilingual,
    parse_reference,
    prepare_bilingual_deck,
    tamil_reference,
    tamil_text,
)
from praisonaippt.core import create_presentation
from praisonaippt.exceptions import SchemaError
from praisonaippt.text_measure import split_balanced
from praisonaippt.utils import split_long_text

ROM_10_17 = "ஆதலால் விசுவாசம் கேள்வியினாலே வரும், கேள்வி தேவனுடைய வசனத்தினாலே வரும்."
BIBLE = {
    "44:10:6": "விசுவாசத்தினாலாகும் நீதியானது: பரலோகத்துக்கு ஏறுகிறவன் யார்?",
    "44:10:7": "அல்லது பாதாளத்துக்கு இறங்குகிறவன் யார்? என்று சொல்லுகிறதுமன்றி;",
    "44:10:17": ROM_10_17,
    "44:1:32": "முதல் அதிகாரத்தின் கடைசி வசனம்.",
    "44:2:1": "இரண்டாம் அதிகாரத்தின் முதல் வசனம்.",
    "38:3:8": "மனுஷன் தேவனை வஞ்சிக்கலாமா?",
    "38:3:9": "நீங்கள் சபிக்கப்பட்டவர்கள்.",
    "38:3:10": "தசமபாகங்களையெல்லாம் பண்டசாலையிலே கொண்டுவாருங்கள்.",
    "18:23:1": "கர்த்தர் என் மேய்ப்பராயிருக்கிறார்.",
    "18:23:2": "அவர் என்னைப் புல்லுள்ள இடங்களில் மேய்த்து,",
}


@pytest.fixture
def bible(tmp_path, monkeypatch):
    path = tmp_path / "tamil-ov.json"
    path.write_text(json.dumps(BIBLE, ensure_ascii=False), encoding="utf-8")
    monkeypatch.setenv("PRAISONAI_TAMIL_BIBLE", str(path))
    return BIBLE


def _deck(verses, **extra):
    deck = {"presentation_title": "T", "sections": [{"section": "", "verses": verses}]}
    deck.update(extra)
    return deck


def _tamil_runs(prs):
    for slide in prs.slides:
        for shape in slide.shapes:
            if not shape.has_text_frame:
                continue
            for p in shape.text_frame.paragraphs:
                for r in p.runs:
                    if any("஀" <= ch <= "௿" for ch in r.text):
                        yield r


# ---------------------------------------------------------------- books & references


@pytest.mark.parametrize("name,english", [
    ("Romans", "Romans"), ("Rom.", "Romans"), ("ரோமர்", "Romans"), ("1 Cor", "1 Corinthians"),
    ("I John", "1 John"), ("1John", "1 John"), ("Psalm", "Psalms"), ("Song of Songs", "Song of Solomon"),
    ("Phil", "Philippians"), ("Deut", "Deuteronomy"), ("எபேசியர்", "Ephesians"),
])
def test_lookup_book(name, english):
    assert lookup_book(name).english == english


def test_lookup_book_rejects_ambiguous_and_unknown():
    with pytest.raises(BookNotFound):
        lookup_book("Jo")          # Joshua / Job / Joel / John / Jonah
    with pytest.raises(BookNotFound):
        lookup_book("Hezekiah")


def test_book_index_matches_tamil_bible_order():
    assert lookup_book("Genesis").index == 0
    assert lookup_book("Romans").index == 44
    assert lookup_book("Revelation").index == 65


def test_parse_reference_forms():
    p = parse_reference("Romans 10:6–12 (NKJV)")
    assert (p.book.english, p.spec, p.version) == ("Romans", "10:6-12", "NKJV")
    assert parse_reference("Malachi 3:8–3:10").spec == "3:8-10"
    assert parse_reference("ரோமர் 10:17 (TAOVBSI)").book.english == "Romans"
    assert parse_reference("Psalm 23").verses == ((23, 0),)
    with pytest.raises(BibleLookupError):
        parse_reference("not a reference")


def test_tamil_reference_label():
    assert tamil_reference("Romans 10:6–12 (NKJV)") == "ரோமர் 10:6-12 (TAOVBSI)"
    assert tamil_reference("Malachi 3:8–3:10", None) == "மல்கியா 3:8-10"


def test_tamil_text_single_range_chapter_and_cross_chapter(bible):
    assert tamil_text("Romans 10:17 (NKJV)", bible) == ROM_10_17
    assert tamil_text("Malachi 3:8–3:10", bible).count(" ") >= 2
    assert tamil_text("Malachi 3:8–3:10", bible).startswith(BIBLE["38:3:8"])
    assert tamil_text("Psalm 23", bible) == f'{BIBLE["18:23:1"]} {BIBLE["18:23:2"]}'
    assert tamil_text("Romans 1:32-2:1", bible) == f'{BIBLE["44:1:32"]} {BIBLE["44:2:1"]}'
    with pytest.raises(BibleLookupError):
        tamil_text("Romans 10:18", bible)


# ---------------------------------------------------------------- fill precedence


def test_auto_fill_from_bible(bible):
    deck = _deck([{"reference": "Romans 10:17 (NKJV)", "text": "So then faith comes by hearing."}],
                 bilingual={"language": "ta"})
    report = fill_tamil(deck)
    verse = deck["sections"][0]["verses"][0]
    assert verse["text_ta"] == ROM_10_17
    assert verse["reference_ta"] == "ரோமர் 10:17 (TAOVBSI)"
    assert report.from_bible == ["Romans 10:17 (NKJV)"] and not report.missing


def test_yaml_text_wins_over_bible(bible):
    deck = _deck([{"reference": "Romans 10:17", "text": "x", "text_ta": "என் சொந்த உரை"}], bilingual="ta")
    report = fill_tamil(deck)
    assert deck["sections"][0]["verses"][0]["text_ta"] == "என் சொந்த உரை"
    assert report.from_yaml == ["Romans 10:17"]


def test_reference_ta_drives_the_lookup(bible):
    deck = _deck([{"reference": "Romans 10:6 (NKJV)", "text": "x", "reference_ta": "ரோமர் 10:6-7 (TAOVBSI)"}],
                 bilingual="ta")
    fill_tamil(deck)
    verse = deck["sections"][0]["verses"][0]
    assert verse["text_ta"] == f'{BIBLE["44:10:6"]} {BIBLE["44:10:7"]}'
    assert verse["reference_ta"] == "ரோமர் 10:6-7 (TAOVBSI)"


def test_missing_verse_reported_and_strict_fails(bible):
    deck = _deck([{"reference": "Romans 10:18", "text": "x"}], bilingual={"language": "ta"})
    report = fill_tamil(copy.deepcopy(deck))
    assert report.missing and "Romans 10:18" in report.summary()
    with pytest.raises(SchemaError):
        fill_tamil(copy.deepcopy(deck), BilingualConfig(enabled=True, strict=True))


def test_highlight_checks(bible):
    deck = _deck([
        {"reference": "Romans 10:17", "text": "faith", "highlights": ["faith"], "highlights_ta": ["இல்லாத சொல்"]},
        {"reference": "Romans 10:6", "text": "speaks", "highlights": ["speaks"]},
    ], bilingual="ta")
    report = fill_tamil(deck)
    assert report.highlight_not_found == [("Romans 10:17", "இல்லாத சொல்")]
    assert report.highlights_missing == ["Romans 10:6"]


def test_verse_opt_out_and_non_scripture_slides_untouched(bible):
    deck = _deck([
        {"reference": "Romans 10:17", "text": "x", "bilingual": False},
        {"slide_type": "title_only", "text": "Romans 10:17"},
    ], bilingual="ta")
    report = fill_tamil(deck)
    assert report.skipped == ["Romans 10:17"]
    assert all("text_ta" not in v for v in deck["sections"][0]["verses"])


def test_source_yaml_never_fetches(bible):
    deck = _deck([{"reference": "Romans 10:17", "text": "x"}], bilingual={"language": "ta", "source": "yaml"})
    report = fill_tamil(deck)
    assert report.missing and "text_ta" not in deck["sections"][0]["verses"][0]


def test_legacy_deck_with_text_ta_is_bilingual_without_block():
    deck = _deck([{"reference": "Romans 10:17", "text": "x", "text_ta": ROM_10_17}])
    cfg = BilingualConfig.from_deck(deck)
    assert cfg.enabled and cfg.source == "yaml"
    assert not BilingualConfig.from_deck(_deck([{"reference": "Romans 10:17", "text": "x"}])).enabled


def test_prepare_does_not_mutate_input(bible):
    deck = _deck([{"reference": "Romans 10:17", "text": "x"}], bilingual="ta")
    before = copy.deepcopy(deck)
    out, report = prepare_bilingual_deck(deck)
    assert deck == before and out["sections"][0]["verses"][0]["text_ta"] == ROM_10_17


def test_unsupported_language_rejected():
    with pytest.raises(SchemaError):
        BilingualConfig.from_deck({"bilingual": "hi"}).validate()


def test_freeze_writes_text_and_switches_source(bible):
    deck = _deck([{"reference": "Romans 10:17", "text": "x"}], bilingual={"language": "ta"})
    frozen, report = freeze_bilingual(deck)
    assert frozen["sections"][0]["verses"][0]["text_ta"] == ROM_10_17
    assert frozen["bilingual"]["source"] == "yaml"
    assert not any(k.startswith("_") for k in frozen)


# ---------------------------------------------------------------- rendering


def _cards_deck(verses, primary="tamil", **style):
    return _deck(verses, bilingual={"language": "ta"}, slide_style={
        "background_color": "#0B1424", "bilingual_layout": "cards",
        "bilingual_primary": primary, **style,
    })


def test_cards_layout_tamil_first_fonts_and_bounds(bible, tmp_path):
    out = create_presentation(_cards_deck([{
        "reference": "Romans 10:17 (NKJV)", "text": "So then faith comes by hearing, and hearing by the word of God.",
        "highlights": ["faith comes"], "highlights_ta": ["விசுவாசம்"],
    }]), output_file=str(tmp_path / "c.pptx"))
    prs = Presentation(out)
    assert len(prs.slides) == 2                     # title + one verse slide
    slide = prs.slides[1]
    texts = [s for s in slide.shapes if s.has_text_frame and s.text_frame.text.strip()]
    ta_box = next(s for s in texts if "ரோமர்" in s.text_frame.text)
    en_box = next(s for s in texts if "Romans" in s.text_frame.text)
    assert ta_box.top < en_box.top                  # bilingual_primary: tamil
    for shape in slide.shapes:
        assert shape.top + shape.height <= prs.slide_height
    for run in _tamil_runs(prs):
        rPr = run._r.rPr
        assert rPr.get("lang") == "ta-IN"
        assert rPr.find(qn("a:cs")).get("typeface") == "Noto Serif Tamil"
        assert rPr.find(qn("a:latin")).get("typeface") == "Noto Serif Tamil"
    # reference and verse are separate paragraphs, never a raw newline in a run
    assert all("\n" not in r.text for r in _tamil_runs(prs))
    cards = [s for s in slide.shapes if s.shape_type == 1]
    assert len(cards) == 2
    assert cards[0]._element.spPr.find(qn("a:solidFill")).find(qn("a:srgbClr")).find(qn("a:alpha")) is not None


def test_cards_split_long_verse_in_step(bible, tmp_path):
    long_en = " ".join(["And they went on and on through the land, rejoicing;"] * 9)
    long_ta = " ".join(["அவர்கள் தேசமெங்கும் சந்தோஷமாய்ப் போனார்கள்;"] * 9)
    out = create_presentation(_cards_deck([{"reference": "Romans 10:17", "text": long_en, "text_ta": long_ta}]),
                              output_file=str(tmp_path / "s.pptx"))
    prs = Presentation(out)
    verse_slides = list(prs.slides)[1:]
    n = len(verse_slides)
    assert n >= 2
    for i, slide in enumerate(verse_slides, 1):
        joined = " ".join(s.text_frame.text for s in slide.shapes if s.has_text_frame)
        assert f"({i}/{n})" in joined
        assert joined.count(f"({i}/{n})") == 2      # both languages carry the part label
        for shape in slide.shapes:
            assert shape.top + shape.height <= prs.slide_height


def test_plan_count_matches_render(bible, tmp_path):
    from praisonaippt.video_exporter import iter_slide_plan

    long_en = " ".join(["And they went on and on through the land, rejoicing;"] * 9)
    deck = _cards_deck([{"reference": "Romans 10:17", "text": long_en}, {"reference": "Romans 10:6", "text": "x"}])
    out = create_presentation(copy.deepcopy(deck), output_file=str(tmp_path / "p.pptx"))
    assert len(list(iter_slide_plan(deck))) == len(Presentation(out).slides)


def test_paired_layout_two_slides_in_primary_order(bible, tmp_path):
    deck = _cards_deck([{"reference": "Romans 10:17 (NKJV)", "text": "So then faith comes by hearing."}],
                       bilingual_layout="paired")
    prs = Presentation(create_presentation(deck, output_file=str(tmp_path / "pr.pptx")))
    assert len(prs.slides) == 3
    first = " ".join(s.text_frame.text for s in prs.slides[1].shapes if s.has_text_frame)
    second = " ".join(s.text_frame.text for s in prs.slides[2].shapes if s.has_text_frame)
    assert "ரோமர் 10:17" in first and "Romans 10:17" in second


def test_below_layout_keeps_bottom_reference(tmp_path):
    deck = _deck([{"reference": "Romans 10:17 (NKJV)", "text": "So then faith comes by hearing.",
                   "reference_ta": "ரோமர் 10:17 (TAOVBSI)", "text_ta": ROM_10_17}],
                 slide_style={"bilingual_layout": "below", "reference_position": "bottom"})
    prs = Presentation(create_presentation(deck, output_file=str(tmp_path / "b.pptx")))
    texts = [s.text_frame.text for s in prs.slides[1].shapes if s.has_text_frame]
    assert any(t.strip() == "Romans 10:17 (NKJV)" for t in texts)     # was silently dropped
    ta_box = next(s for s in prs.slides[1].shapes if s.has_text_frame and "ரோமர்" in s.text_frame.text)
    assert [p.text for p in ta_box.text_frame.paragraphs] == ["ரோமர் 10:17 (TAOVBSI)", ROM_10_17]


def test_section_tamil_only_in_bilingual_decks(bible, tmp_path):
    section = {"section": "1. Sent", "section_ta": "1. அனுப்பப்படுதல்",
               "verses": [{"reference": "Romans 10:17", "text": "x"}]}
    en = {"presentation_title": "T", "sections": [copy.deepcopy(section)]}
    prs = Presentation(create_presentation(en, output_file=str(tmp_path / "en.pptx")))
    assert "அனுப்பப்படுதல்" not in prs.slides[1].shapes[0].text_frame.text
    ta = dict(en, bilingual="ta")
    prs = Presentation(create_presentation(ta, output_file=str(tmp_path / "ta.pptx")))
    assert any("அனுப்பப்படுதல்" in s.text_frame.text for s in prs.slides[1].shapes if s.has_text_frame)


def test_mixed_title_run_split_by_script(bible, tmp_path):
    deck = _cards_deck([{"reference": "Romans 10:17", "text": "x"}])
    deck["presentation_title"] = "He will not put to shame / அவன் வெட்கப்படுவதில்லை"
    prs = Presentation(create_presentation(deck, output_file=str(tmp_path / "t.pptx")))
    runs = [r for s in prs.slides[0].shapes if s.has_text_frame for p in s.text_frame.paragraphs for r in p.runs]
    english = next(r for r in runs if "shame" in r.text)
    tamil = next(r for r in runs if "அவன்" in r.text)
    assert english.font.name != "Noto Serif Tamil"
    assert tamil._r.rPr.find(qn("a:cs")).get("typeface") == "Noto Serif Tamil"


def test_background_image_is_cropped_not_stretched(tmp_path):
    from PIL import Image

    img = tmp_path / "square.png"
    Image.new("RGB", (400, 400), "navy").save(img)
    deck = _deck([{"reference": "Romans 10:17", "text": "x"}], slide_style={"background_image": str(img)})
    prs = Presentation(create_presentation(deck, output_file=str(tmp_path / "bg.pptx")))
    pic = next(s for s in prs.slides[1].shapes if s.shape_type == 13)
    assert pic.crop_top > 0 and pic.crop_bottom > 0 and pic.crop_left == 0
    deck["slide_style"]["background_fit"] = "stretch"
    prs = Presentation(create_presentation(deck, output_file=str(tmp_path / "bg2.pptx")))
    pic = next(s for s in prs.slides[1].shapes if s.shape_type == 13)
    assert pic.crop_top == 0


# ---------------------------------------------------------------- splitting, loader, CLI


def test_split_balanced_never_cuts_a_highlight():
    text = "Bring all the tithes into the storehouse, That there may be food in My house, And try Me now in this."
    parts = split_balanced(text, 2, ["food in My house"])
    assert len(parts) == 2 and any("food in My house" in p for p in parts)


def test_split_long_text_clause_fallback_is_opt_in():
    sentence = "அவர் வந்தார், அவர் பார்த்தார், அவர் பேசினார், அவர் போனார், அவர் திரும்பினார்"
    assert split_long_text(sentence, max_length=50) == [sentence]
    parts = split_long_text(sentence, max_length=50, clause_fallback=True)
    assert len(parts) == 2 and " ".join(parts) == sentence


def test_sections_from(tmp_path):
    from praisonaippt.loader import load_verses_from_file

    (tmp_path / "en.yaml").write_text(
        "presentation_title: E\nsections:\n- section: ''\n  verses:\n  - reference: Romans 10:17\n    text: x\n",
        encoding="utf-8")
    (tmp_path / "ta.yaml").write_text("presentation_title: T\nsections_from: en.yaml\n", encoding="utf-8")
    data = load_verses_from_file(str(tmp_path / "ta.yaml"))
    assert data["presentation_title"] == "T"
    assert data["sections"][0]["verses"][0]["reference"] == "Romans 10:17"
    (tmp_path / "bad.yaml").write_text("sections_from: en.yaml\nsections: []\n", encoding="utf-8")
    assert load_verses_from_file(str(tmp_path / "bad.yaml")) is not None   # empty sections allowed
    (tmp_path / "both.yaml").write_text(
        "sections_from: en.yaml\nsections:\n- section: A\n", encoding="utf-8")
    assert load_verses_from_file(str(tmp_path / "both.yaml")) is None


def test_bilingual_fill_cli(bible, tmp_path):
    import yaml

    from praisonaippt.cli import handle_bilingual_fill_command

    src = tmp_path / "deck.yaml"
    src.write_text("presentation_title: T\nbilingual: ta\nsections:\n- section: ''\n  verses:\n"
                   "  - reference: Romans 10:17 (NKJV)\n    text: x\n", encoding="utf-8")
    out = tmp_path / "filled.yaml"
    args = Namespace(input_file=str(src), input=None, output=str(out), bilingual=None)
    assert handle_bilingual_fill_command(args) == 0
    filled = yaml.safe_load(out.read_text(encoding="utf-8"))
    assert filled["sections"][0]["verses"][0]["text_ta"] == ROM_10_17
